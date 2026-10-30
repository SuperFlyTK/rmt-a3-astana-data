"""Check the report against its archived inputs and the Assignment 4 minima."""
from pathlib import Path
from zipfile import ZipFile
import csv
import hashlib
import json
import re
import subprocess
import xml.etree.ElementTree as ET
from decimal import Decimal
from datetime import datetime

from docx import Document
from build_report import fmt

ROOT = Path(__file__).resolve().parents[1]
HERE = ROOT / "assignment4"


def read(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def verify():
    counts = json.loads((HERE / "submission_checks.json").read_text())
    manifest = json.loads((HERE / "evidence_manifest.json").read_text())
    report = ROOT / counts["report"]
    checks = {}
    def check(name, condition):
        checks[name] = bool(condition)
        if not condition:
            raise AssertionError(name)
    check("abstract_at_most_150", counts["abstract_words"] <= 150)
    check("results_at_least_800", counts["results_discussion_words"] >= 800)
    check("conclusion_at_least_300", counts["conclusion_words"] >= 300)
    check("block_e_interpretation_at_least_300", counts["block_e_reliability_words"] >= 300)
    check("at_least_five_graphs", counts["results_figures"] >= 5)
    check("at_least_three_results_tables", counts["results_tables"] >= 3)
    check("report_hash", hashlib.sha256(report.read_bytes()).hexdigest() == counts["report_sha256"])
    check("decimal_rounding_contrast", fmt(271.17499999999995)=="271.18")
    check("decimal_rounding_iqr", fmt(230.2249999999999)=="230.23")
    for category in ("inputs", "scripts", "outputs"):
        for path, digest in manifest[category].items():
            payload=(ROOT/path).read_bytes()
            if Path(path).suffix in {".csv", ".json", ".md", ".py", ".txt"}:
                payload=payload.replace(b"\r\n",b"\n")
            check("hash:"+path, hashlib.sha256(payload).hexdigest() == digest)
    for fig in manifest["figures"]:
        check("hash:"+fig["file"], hashlib.sha256((ROOT/fig["file"]).read_bytes()).hexdigest() == fig["sha256"])
    for name in ("docs/expected_results.md", "docs/expected_results.json"):
        lookup = subprocess.run(["git", "show", "b4614355903efdc800f870d2678d2416e0490f81:"+name],cwd=ROOT,capture_output=True)
        current=(ROOT/name).read_bytes().replace(b"\r\n",b"\n")
        if lookup.returncode==0:
            unchanged=lookup.stdout.replace(b"\r\n",b"\n")==current
        else:
            origin=json.loads((ROOT/'docs/expected_results_origin.json').read_text())
            unchanged=hashlib.sha256(current).hexdigest()==origin['normalized_sha256'][name]
        check("preregistered_unchanged:"+name,unchanged)
    doc = Document(report)
    text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    check("correct_first_line", doc.paragraphs[0].text.startswith("Group SE-2424 | Assignment 4"))
    check("exact_a3_question", counts["rq"] in text)
    check("approach_card_unit", "Unit of analysis" in table_text)
    check("approach_card_rejected_alternative", "Alternative rejected" in table_text)
    expected = read(HERE / "tables/expectations.csv")
    check("all_eight_expectations", len(expected) == 8)
    for row in expected:
        check("verbatim:"+row["id"], row["expected_result_unchanged"] in table_text)
    check("all_five_objectives", len(read(HERE / "tables/objectives.csv")) == 5)
    check("six_embedded_figures", len(doc.inline_shapes) == counts["results_figures"])
    # Every data table and graph has at least two complete interpretation sentences.
    children = list(doc.element.body)
    for i,el in enumerate(children):
        if el.tag.endswith("}p"):
            label="".join(el.itertext())
            if not re.match(r"(?:Table|Figure) \d+\.",label):
                continue
            sentences=[]
            for following in children[i+1:]:
                if following.tag.endswith("}tbl"):
                    continue
                txt="".join(following.itertext())
                style=following.find(".//"+"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle")
                if sentences and (re.match(r"(?:Table|Figure) \d+\.",txt) or (style is not None and "Heading" in str(style.attrib))):
                    break
                if txt.startswith("Evidence:") or not txt.strip():
                    continue
                sentences.extend(re.findall(r"[^.!?]+[.!?](?:\s|$)",txt))
                if len(sentences)>=2:
                    break
            check("interpretation:"+label[:90],len(sentences)>=2)
    links = json.loads((HERE / "report_links.json").read_text())
    check("no_zip_cloud_links", not any("drive.google.com/file/" in url or url.endswith(".zip") for url in links))
    rows=read(HERE/"data/ors_processed.csv")
    check("all_ors_rows_retained", len(rows)==572 and all(r["usable"]=="true" for r in rows))
    batches=read(HERE/"tables/batch_results.csv")
    check("four_complete_followup_windows", len([r for r in batches if r["phase"]=="followup" and r["n_valid_pairs"]=="20"])==4)
    original_rows=[]
    for path in [*sorted((ROOT/"data/raw").glob("*.csv")),ROOT/"data/exploratory/ors/observations.csv",ROOT/"data/followup/ors/observations.csv"]:
        original_rows.extend(read(path))
    matches=0
    for row in original_rows:
        saved=json.loads((ROOT/row["raw_file"]).read_text(encoding="utf-8-sig"))["features"][0]["properties"]["summary"]
        if all(Decimal(row[field])==Decimal(str(saved[key])) for field,key in [("duration_s","duration"),("distance_m","distance")]):
            matches+=1
    check("every_ors_measurement_matches_original_response",matches==len(original_rows)==572)
    google=read(ROOT/"data/followup/google_request_status.csv")
    for batch in (r for r in batches if r["phase"]=="followup"):
        cid=batch["cycle_id"]
        g={r["route_id"]:r for r in google if r["cycle_id"]==cid}
        o={r["route_id"]:r for r in original_rows if r["cycle_id"]==cid}
        check("matched_routes:"+cid,len(g)==len(o)==len(g.keys()&o.keys())==20)
        actual=max(abs((datetime.fromisoformat(g[k]["timestamp_utc"])-datetime.fromisoformat(o[k]["timestamp_utc"])).total_seconds()) for k in g)
        check("saved_pair_timing:"+cid,abs(actual-float(batch["max_request_start_gap_s"]))<1e-9 and actual<=60)
    result=dict(checks=checks,passed=len(checks),failed=0)
    (HERE/"verification.json").write_text(json.dumps(result,indent=2),encoding="utf-8")
    print(json.dumps(dict(passed=len(checks),failed=0)))


if __name__ == "__main__":
    verify()
