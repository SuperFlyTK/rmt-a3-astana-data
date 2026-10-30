"""Create the Assignment 4 manuscript from the checked evidence tables."""
from __future__ import annotations

import csv
import hashlib
import json
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
HERE = ROOT / "assignment4"
TABLES = HERE / "tables"
TEAM = "AkhanovYerdaulet_ShadiyarbekSyrym_SabyrzhanAkylbek"
OUTPUT = HERE / f"RMT_A4_Team_{TEAM}.docx"
BASE = "https://github.com/SuperFlyTK/rmt-a3-astana-data"
DATA_BASE = "https://github.com/SuperFlyTK/rmt-a3-astana-data"
RQ = "How do Google traffic-aware travel-time estimates differ from ORS control estimates on fixed Astana routes during peak and off-peak periods?"


def read(name):
    with (TABLES / name).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def words(text):
    return len(re.findall(r"\b[\w]+(?:[-'][\w]+)*\b", text))


def text_hash(path):
    return hashlib.sha256(path.read_bytes().replace(b"\r\n",b"\n")).hexdigest()


def fmt(value, digits=2):
    # Remove binary storage noise below a microsecond before decimal rounding.
    return str(Decimal(str(round(float(value),6))).quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP))


def build():
    inv = json.loads((TABLES / "inventory.json").read_text())
    summary = json.loads((TABLES / "period_summary.json").read_text())
    batches = read("batch_results.csv")
    follow = [r for r in batches if r["phase"] == "followup"]
    groups = {r["period"]:r for r in summary["groups"]}
    peak = groups["peak"]["median_of_window_median_absolute_differences_s"]
    off = groups["offpeak"]["median_of_window_median_absolute_differences_s"]
    gap = summary["peak_minus_offpeak_window_median_s"]
    ratio = summary["peak_to_offpeak_window_median_ratio"]
    assert inv["ors_observations"] == 572 and summary["valid_windows"] == 4, "Review prose when changing the study snapshot."
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for name, size in [("Title",19),("Heading 1",14),("Heading 2",12),("Heading 3",11),("Caption",10)]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0,0,0)
        style.font.underline = False
        fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
        if fonts is not None:
            for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                fonts.attrib.pop(qn("w:"+key), None)
        style.paragraph_format.space_before = Pt(10 if name.startswith("Heading") else 3)
        style.paragraph_format.space_after = Pt(5)
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
    doc.core_properties.author = "Yerdaulet Akhanov; Syrym Shadiyarbek; Akylbek Sabyrzhan"
    doc.core_properties.title = "Peak and Off Peak Differences Between Google and ORS Travel Time Estimates on Astana Routes"
    doc.core_properties.subject = "RMT Assignment 4, SE-2424"
    doc.core_properties.comments = ""
    body, abstract_text, conclusion_text = [], [], []
    segment = "front"
    table_count, figure_count = 0, 0
    links = set()
    table_interpretations, figure_interpretations = [], []

    def p(text, style=None):
        para = doc.add_paragraph(text, style)
        para.paragraph_format.widow_control = True
        if segment == "results" and style is None:
            body.append(text)
        elif segment == "conclusion" and style is None:
            conclusion_text.append(text)
        return para

    def h(text, level=2):
        doc.add_heading(text, level)

    def link(para, label, url):
        links.add(url)
        element = OxmlElement("w:hyperlink")
        rid = para.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        element.set(qn("r:id"),rid)
        run = OxmlElement("w:r")
        props = OxmlElement("w:rPr")
        color = OxmlElement("w:color"); color.set(qn("w:val"),"174A70"); props.append(color)
        under = OxmlElement("w:u"); under.set(qn("w:val"),"single"); props.append(under)
        size = OxmlElement("w:sz");size.set(qn("w:val"),"19");props.append(size)
        run.append(props)
        text = OxmlElement("w:t");text.text=label;run.append(text);element.append(run)
        para._p.append(element)

    def evidence(*paths):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(5)
        para.add_run("Evidence: ").font.size=Pt(9.5)
        for i,path in enumerate(paths):
            if i:para.add_run("; ")
            path = {"docs/research_log.md":"docs/collection_log.csv", "docs/api_check.md":"docs/request_parameters.json"}.get(path,path)
            link(para,path,BASE+"/blob/main/research/"+path)
        return para

    def table(title, headers, rows, widths=None, sources=(), interpretations=(), numbered=True):
        nonlocal table_count
        if numbered:
            table_count += 1
            caption = f"Table {table_count}. {title}"
        else:caption=title
        cap = doc.add_paragraph(caption,"Caption")
        cap.paragraph_format.keep_with_next=True
        t=doc.add_table(rows=1, cols=len(headers))
        t.alignment=WD_TABLE_ALIGNMENT.CENTER
        t.autofit=False
        widths=widths or [17/len(headers)]*len(headers)
        for col,width in zip(t.columns,widths):col.width=Cm(width)
        for c,width,text in zip(t.rows[0].cells,widths,headers):
            c.width=Cm(width);c.text=text
            shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"EAEAEA");c._tc.get_or_add_tcPr().append(shade)
        for row in rows:
            for c,width,text in zip(t.add_row().cells,widths,row):
                c.width=Cm(width);c.text=str(text)
        for i,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if i==0:
                repeat=OxmlElement("w:tblHeader");row._tr.get_or_add_trPr().append(repeat)
            for c in row.cells:
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=c._tc.get_or_add_tcPr()
                margins=OxmlElement("w:tcMar")
                for edge in ("top","left","bottom","right"):
                    el=OxmlElement("w:"+edge);el.set(qn("w:w"),"65");el.set(qn("w:type"),"dxa");margins.append(el)
                tcpr.append(margins)
                for para in c.paragraphs:
                    para.paragraph_format.space_after=Pt(3)
                    para.paragraph_format.space_before=Pt(2)
                    para.paragraph_format.line_spacing=1.02
                    # Avoid a header with one isolated row, or one final orphan row.
                    para.paragraph_format.keep_with_next = i < 2 or i == len(t.rows)-2
                    for run in para.runs:
                        run.font.size=Pt(9.5)
                        run.bold=i==0
        borders=OxmlElement("w:tblBorders")
        for edge in ("top","left","bottom","right","insideH","insideV"):
            el=OxmlElement("w:"+edge);el.set(qn("w:val"),"single");el.set(qn("w:sz"),"4");el.set(qn("w:color"),"D9D9D9");borders.append(el)
        t._tbl.tblPr.append(borders)
        if sources:evidence(*sources)
        for text in interpretations:p(text)
        if numbered:table_interpretations.append(dict(number=table_count, paragraphs=len(interpretations)))
        return t

    def figure(name, caption, interpretations, sources):
        nonlocal figure_count
        figure_count+=1
        para=doc.add_paragraph();para.paragraph_format.keep_with_next=True
        run=para.add_run();run.add_picture(str(HERE/"figures"/name),width=Cm(16.8))
        inline=run._r.xpath(".//wp:docPr")[0];inline.set("descr",caption)
        cap=doc.add_paragraph(f"Figure {figure_count}. {caption}","Caption");cap.paragraph_format.keep_with_next=True
        evidence(*sources)
        for text in interpretations:p(text)
        figure_interpretations.append(dict(number=figure_count,paragraphs=len(interpretations)))

    p("Group SE-2424 | Assignment 4 | Akhanov Yerdaulet, Shadiyarbek Syrym, Sabyrzhan Akylbek")
    doc.add_paragraph(doc.core_properties.title,"Title")
    para=p("Public research folder: ");link(para,"Browse the research files",BASE+"/tree/main/research/assignment4")
    para=p("Public source records: ");link(para,"Collection records",BASE+"/tree/main/research/data")
    p("Research question: "+RQ)
    count_para=p("Results and Discussion count is verified at build time.")
    h("Abstract")
    abstract=(f"Navigation services can return different travel times for the same endpoints. We compared Google traffic-aware estimates with openrouteservice (ORS), which supplied a control without live traffic, on 20 directed Astana routes. Archived evidence covers two dates and four planned comparison windows, with 80 valid pairs. The median of window median absolute differences was {fmt(peak)} seconds during peak periods and {fmt(off)} seconds off peak, a descriptive gap of {fmt(gap)} seconds. All 572 archived ORS observations passed the recorded checks, but ORS durations stayed unchanged within each route. The original scheduled collection target was missed. Individual Google estimates were not retained, preventing route-level reanalysis and significance testing. These findings describe disagreement for the sampled routes and times; they establish neither prediction accuracy nor a causal traffic effect.")
    p(abstract);abstract_text.append(abstract)
    p(f"Abstract word count: {words(abstract)}")
    p("Keywords: travel time; provider disagreement; Google Routes; openrouteservice; peak periods; Astana")
    p("The title names the measured difference, the routing services, the sampled routes and the city. It keeps the peak-period variable visible without suggesting that we measured actual driving accuracy.")
    h("Paper structure",1)
    table("Article structure",["Section","Subsections and role"],[
        ["1 Introduction","1.1 Background and problem context; 1.2 literature, disagreement and source selection; 1.3 gap; 1.4 research question and aim; 1.5 objectives"],
        ["2 Methodology","2.1 Approach Card; 2.2 source and sample; 2.3 collection and preparation; 2.4 metrics and interpretation; 2.5 ethics and provenance"],
        ["3 Results and Discussion","3.1 inventory; 3.2 preparation and variables; 3.3 pipeline; 3.4 window results; 3.5 objectives; 3.6 expected versus observed; 3.7 Block E and reliability; 3.8 sensitivity; 3.9 additional interpretation; 3.10 literature comparison; 3.11 validity"],
        ["4 Conclusion","Answer, contribution, unresolved claims and specific next steps"],
        ["References and appendices","Harvard references with full text; A evidence and reproducibility; B instrument provenance and original scenarios"]
    ],[3.5,13.5],numbered=False)
    p("The introduction explains why provider disagreement matters, and the methods define what can be measured from the archived evidence. The results then answer the amended question while keeping the original objectives and expectations visible.")

    h("1 Introduction",1)
    h("1 1 Background and problem context")
    p("Drivers and transport researchers use navigation estimates to judge journey duration. Two services can recommend different streets, assign different speeds to the same streets, or react differently to departure conditions. A difference between their estimates therefore needs explanation before either estimate is treated as a reliable description of a trip in Astana.")
    p("Three concepts guide this study. Agreement describes how close two services are to each other. Variability describes how their difference changes between sampled routes or times. Accuracy concerns the difference between a prediction and an independently observed journey. We collected service outputs without observing car journeys, so the evidence supports agreement analysis.")
    h("1 2 Literature review and disagreement")
    p("Derrow-Pinion et al. (2021) describe a Google Maps prediction system that uses road structure and traffic information. Soon et al. (2023) review traffic forecasting and show why temporal coverage, input quality and unusual events matter. These papers explain the value of matching endpoints and departure conditions. They do not establish the internal design of the services used in our present comparison.")
    p("Fu et al. (2023) compare routing methods for common origin-destination pairs. Verma et al. (2025) show that strong statistical fit between routing outputs can coexist with differences in average duration. Boeing and Zhou (2026) develop a model for minimally congested travel using Google estimates as its validation target. Together, these studies support examining differences in time units and stating clearly what the reference output represents.")
    p("The literature also contains a meaningful divergence in reported agreement. Elnaggar et al. (2025) report close agreement in a comparison involving Google navigation estimates, whereas Bastarianto et al. (2025) and Valero et al. (2025) find mismatches between API estimates and reported journeys. The comparison targets differ: agreement between service outputs does not establish agreement with a traveller's experience, and self-reported journeys introduce recall and rounding uncertainty. Our evidence can examine the first issue but cannot resolve the second.")
    p("Mansurova et al. (2025) provide empirical bus operations data from Astana, separating running and stopping behaviour. This establishes useful local context, but bus operations cannot validate private-car estimates directly. Hodson (2022) explains why absolute and squared differences have different purposes. We use the median to describe a typical gap and squared differences to make larger gaps visible, without calling either an error against observed travel.")
    h("1 2 1 Review scope and source selection",3)
    p("The Assignment 2 review retained ten scholarly sources within its declared five-year eligibility window, from 3 October 2021 to the search date of 3 October 2026. Searches combined broad travel-time forecasting queries, focused routing-service comparisons and critical searches about validation and reference outcomes. Inclusion required relevant subject matter, identifiable bibliographic details and accessible full text. Excluded records remained separate from the retained corpus. The present manuscript carries forward that review and checks the retained full-text links again; it does not claim a new systematic search.")
    h("1 3 Problem statement and research gap")
    p("The Assignment 2 literature did not resolve how traffic-aware and open-routing control estimates differ on these Astana routes across the predefined periods. A local comparison can document that disagreement, but it cannot separate traffic inputs from provider-specific routing and speed assumptions. The initial provider designs were constrained by access and publication permissions, leading to the dated Google-ORS amendment recorded in Assignment 3.")
    h("1 4 Research question and aim")
    p(RQ)
    p("The question is identical to the submitted Assignment 3 wording. The aim is to describe the size, direction and observed period variation of disagreement between Google and the ORS control, while reporting collection quality and the limits of the retained evidence.")
    h("1 5 Objectives carried forward")
    p("The five objectives originally declared in Assignment 1 remain the accountability framework. They include a longer observation period, different providers and inferential analyses that the completed study could not deliver. Table 5 quotes each objective and records its actual status; the amended descriptive comparison does not make those earlier promises disappear.")

    h("2 Methodology",1)
    h("2 1 Approach Card")
    table("Approach Card",["Dimension","Declared approach"],[
        ["Approach and design","Quantitative secondary-data analysis; observational repeated measurements on fixed directed routes."],
        ["Unit of analysis","A provider response is a recorded observation. The retained Google-ORS period comparison uses a window summary; repeated requests are not independent city samples."],
        ["Research strategy","Bounded Astana case study with purposively selected public endpoints and predefined time categories."],
        ["Alternative rejected","A causal traffic experiment was rejected because traffic and provider models were not controlled. An accuracy study was rejected because no independently observed car journeys were collected."],
        ["Strength","Stable route identifiers, close request starts, preserved ORS records and dated protocols support inspection."],
        ["Weakness","Only two observation dates, dependent routes, an unchanged ORS control and aggregate-only Google retention restrict inference."]
    ],[3.5,13.5],numbered=False)
    p("This card includes the unit of analysis and the rejected alternative requested in the Assignment 3 feedback. The window is the available comparison unit because individual Google estimates were not archived.")
    h("2 2 Sources and sample")
    p("The route list fixes public OpenStreetMap road nodes near landmarks, with each direction treated separately. The archived Overpass query, retrieval metadata and selected endpoints identify the source and selection process. Road-node coordinates describe endpoint vicinities and do not guarantee building entrances. Straight-line distances in the route inventory are geographic lower bounds, not measured driving distances.")
    evidence("data/routes.csv","data/sources/astana_overpass_query.txt","data/sources/astana_osm_manifest.json")
    p("The local period definitions are unchanged from Assignment 3: peak is 17:30-19:30, while off-peak is 10:00-12:00 or 22:00-24:00, in Astana time, UTC+05:00. Each interval includes its start and excludes its end. Everything outside these intervals is labelled other. The original schedule, revised cutoff and actual windows remain distinguishable in the protocol files.")
    evidence("config/study.json","docs/google_followup_protocol.json","docs/submission_deadline.json")
    h("2 3 Collection and preparation")
    p("The provider choices changed across the project: Assignment 1 proposed Google Maps and Yandex Maps, Assignment 2 reviewed Google Maps and 2GIS, and the initial Assignment 3 registration named TomTom and 2GIS. The dated Assignment 3 amendment established Google and ORS as the active comparison. We retain the earlier objectives and expectations with their original provider names, while answering the amended Assignment 3 question.")
    p("ORS observations retain the request start, route and cycle identifiers, returned duration and distance, HTTP outcome and latency. The paired collection queried Google with TRAFFIC_AWARE_OPTIMAL and ORS for identical endpoints, allowing a request-start separation of at most 60 seconds. Individual Google durations, distances and geometry were processed transiently; only the recorded comparison aggregates and request metadata remain. This retention decision is a limitation of replication.")
    p("The updated evidence includes all collection phases archived by 25 October 2026 at 17:30:53 Astana time. It adds the final scheduled records and final paired window that were absent from the earlier public export. Assignment 4 introduces an offline evidence-building script, coverage and missingness tables, sensitivity checks and six plots. It makes no new routing requests, changes no archived measurement and does not backfill missing times.")
    evidence("assignment4/evidence_manifest.json","docs/research_log.md","assignment4/analyse.py")
    h("2 4 Metrics and interpretation rules")
    p("For a retained pair, let G and O denote Google and ORS durations in seconds. The signed difference is d=G-O, and the absolute difference is |d|. A positive signed difference means Google returned a longer estimate. MAE is the mean of |d|; RMSE is the square root of the mean of d squared. The mean symmetric absolute percentage difference is the mean of 100|G-O|/((G+O)/2). The IQR is the 75th percentile minus the 25th percentile, using linearly interpolated quartiles.")
    p("The period summary is the median of its recorded window medians. It is not the median of all individual pairs, which cannot be reconstructed from the archived Google evidence. The peak-off-peak contrast subtracts the off-peak summary from the peak summary; their ratio divides peak by off peak. Length comparability requires 100|LG-LO|/((LG+LO)/2) at most five percent. Similar length cannot establish that two routes use the same streets.")
    p("Reported time summaries are rounded to two decimal places, while the evidence files retain their original precision. Period contrasts are calculated before rounding the displayed period values.")
    p("The original route-by-day test and confidence intervals are withheld because the retained window summaries do not recover paired route-level differences. We report descriptive contrasts and show how the result changes when a window is omitted. No trained model, training-test split or causal treatment is used; the selected method remains Block E from Assignment 3.")
    h("2 5 Ethics and data provenance")
    p("The study uses public endpoints and service outputs, with no recruited participants or personal journey histories. Credentials are excluded from the public files. ORS material is attributed to openrouteservice by HeiGIT and OpenStreetMap, and the endpoint source is attributed to OpenStreetMap contributors. The submission evidence folders are directly browsable, replacing the ZIP-only access criticised in the feedback. The original expected-results files and frozen Assignment 3 snapshots are preserved in the collection archive.")
    evidence("data/sources/README.md","docs/api_check.md","docs/expected_results.md")

    segment="results"
    h("3 Results and Discussion",1)
    h("3 1 Inventory of collected evidence")
    table("Collection inventory",["Evidence","Target and achieved","Form and location","Analysis"],[
        ["Directed routes","20 fixed; 20 retained","data/routes.csv","Route identity and coverage"],
        ["Main ORS series","60 pilot + 412 labelled scheduled = 472","assignment4/data/ors_observations.csv","Collection quality; 4 of the 412 were manual verification"],
        ["Additional ORS","20 exploratory + 80 follow-up = 100","Same file; phase labels retained","Paired collection checks; exploratory cycle separate"],
        ["Google-ORS follow-up","80 planned pairs; 80 valid in 4/4 windows","google_ors_summary.csv and collection_windows.csv in assignment4/data","Period disagreement"],
        ["Google metadata","100 statuses: 20 exploratory + 80 follow-up","assignment4/data/google_request_status.csv","Request outcomes and coverage"],
        ["Original automatic target","768 nominal; 408 automatic attempts, 53.125%","assignment4/tables/inventory.json; docs/cloud_runs.json","ER8; no backfill"],
        ["Google individual outputs","0 retained","Aggregate-only retention in the protocol","Route-level reanalysis unavailable"]
    ],[2.7,4.7,5.8,3.8],sources=("assignment4/tables/inventory.json","assignment4/data/collection_windows.csv","docs/cloud_runs.json"),interpretations=[
        "The final archive holds 572 ORS observations, while the planned comparison contains 80 valid pairs across four windows. Keeping pilot, scheduled, exploratory and follow-up labels prevents instrument checks from inflating the comparison sample. The collection achieved the amended paired-window target but missed the original automatic collection target."])
    p("The source records span 24-25 October 2026. The earlier public export held 508 ORS rows and only three follow-up windows; the final archive adds 64 ORS rows, 20 Google request statuses and one comparison summary. These are later archived observations, not new data collected for this assignment.")
    figure("03_coverage.png","ORS coverage by collection phase and local time category.",[
        "Most main-series observations fell outside the comparison periods: 372 of 472 were other, leaving 36 peak and 64 off-peak records. Across all phases there were 76 peak, 104 off-peak and 392 other ORS records. The exploratory cycle belongs entirely to other, so including it in the period contrast would give the comparison a misleading sample definition."
    ],("assignment4/tables/coverage.csv","assignment4/tables/inventory.json"))

    h("3 2 Data preparation and measurement levels")
    cleaning=read("cleaning.csv")
    table("Preparation record",["Operation","Rule applied","Affected and retained"],
          [[r["operation"],r["rule"],f"Input {r['input_rows']}; affected {r['affected']}; output {r['output_rows']}"] for r in cleaning],
          [4.2,8.6,4.2],sources=("assignment4/tables/cleaning.csv","assignment4/data/ors_processed.csv"),interpretations=[
              "All 572 ORS rows pass the recorded screen, so no raw observation was deleted or replaced. Removing other from a particular period comparison is an analysis filter, while the records remain in the archive. Removing the exploratory summary from follow-up analysis similarly preserves its evidence while respecting its separate purpose."])
    table("Variables used and missing values",["Variable or field group","Level and summary","Missing out of 572"],[
        ["duration_s","Ratio; min 510.5, median 899.5, max 1578.8 seconds","0"],
        ["distance_m","Ratio; min 4254.9, median 11002.0, max 20952.3 metres","0"],
        ["latency_s","Ratio; min 0.296, median 0.585, max 5.446 seconds","0"],
        ["speed_kmh","Ratio, derived; min 30.01, median 42.57, max 53.80 km/h","0"],
        ["duration_no_traffic_s","Ratio; ORS supplies no separate value","572 (100%)"],
        ["route_id; cycle_id; phase; provider","Nominal identifiers and categories; distributions in coverage.csv","0 in each"],
        ["timestamp_utc; timestamp_astana","Interval time; timezone-aware request starts","0 in each"],
        ["status; http_code; error","Nominal; status ok and HTTP 200 throughout; error blank on success","error: 572 (100%); status/http_code: 0"],
        ["period; date_astana; usable","Derived nominal period/status and interval date; full counts retained","0 in each"]
    ],[4.7,8.3,4],sources=("assignment4/tables/variables.csv","assignment4/data/data_dictionary.csv"),interpretations=[
        "The absent no-traffic field cannot be interpreted as zero travel time or used to calculate congestion delay. Empty error messages are expected for successful requests and have a different meaning from missing measurements. Duration, distance and latency are complete, which supports the stated summaries without imputation."])
    p(f"The overall blank-cell count in the raw ORS export is {inv['raw_blank_cells']}/{inv['raw_cells']} ({inv['raw_blank_pct']:.2f}%), across its 14 columns. Every blank is in the optional no-traffic or success-error field. For the 12 required observation fields, the count is 0/{inv['required_cells']}; this second denominator makes clear that required measurement completeness and structural blanks answer different questions.")
    evidence("assignment4/tables/inventory.json")

    h("3 3 Analysis pipeline and interpretation steps")
    p("We first join the archived ORS phases without changing their values and check route identities against the fixed inventory. The processing script applies the recorded success, timestamp, speed and duplicate rules, then adds local dates and periods. Duplicate checking uses exact rows, repeated request identifiers and conflicting phase-cycle-route-provider identities; neither repeated equal durations at different timestamps nor reverse routes count as duplicate requests.")
    p("Next, we check that each Google batch agrees with its protocol and archived request-status counts. The exploratory cycle is displayed separately. The four registered follow-up batches all contain 20 valid pairs and satisfy the minimum of five pairs for a numeric summary. Their recorded maximum request-start gaps range from 0.000702 to 0.001219 seconds, comfortably within the 60-second pairing rule. These checks support the batch identities and timing; they cannot reconstruct missing individual Google values.")
    p("For interpretation, we first establish which rows and windows exist, then compare the magnitude and direction of the saved differences. We inspect route-length comparability before discussing a traffic-related explanation. Finally, we compare the observations with each original expectation, examine sensitivity and restrict the answer to the sampled routes and dates. This order prevents a clear-looking chart from being mistaken for evidence of accuracy or causation.")
    evidence("assignment4/analyse.py","assignment4/tables/batch_results.csv","analysis/export_public_data.py")

    h("3 4 Results of the recorded comparison windows")
    names=["24 Oct peak","24 Oct off-peak","25 Oct morning off-peak","25 Oct peak"]
    table("All planned comparison windows",["Window","Valid pairs","Median signed / absolute gap s","MAE / RMSE s","Signed IQR s"],
          [[name,r["n_valid_pairs"],f"{fmt(r['median_signed_difference_s'])} / {fmt(r['median_absolute_difference_s'])}",f"{fmt(r['mae_s'])} / {fmt(r['rmse_s'])}",fmt(r['iqr_signed_difference_s'])] for name,r in zip(names,follow)],
          [4,1.5,4.6,4.3,2.6],sources=("assignment4/tables/batch_results.csv",),interpretations=[
              "All four saved median signed differences are positive, showing that the central Google estimate exceeded the ORS control in each window. The largest median absolute gap occurred in the first peak window, while the smallest occurred in the morning off-peak window. This establishes a pattern among the archived summaries, without showing the size of each individual route's gap."])
    figure("01_window_disagreement.png","Median absolute provider disagreement in each follow-up window.",[
        "The recorded window medians were 748.85 and 610.70 seconds in peak periods, compared with 437.80 and 379.40 seconds off peak. Both peak medians exceed both off-peak medians. However, each window was sampled in a short batch near its recorded start, so these points do not describe every moment of the longer time intervals."
    ],("assignment4/tables/batch_results.csv",))
    figure("02_period_comparison.png","Peak and off-peak summaries with equal weight for each recorded window.",[
        f"The peak summary is {fmt(peak)} seconds ({fmt(peak/60)} minutes), compared with {fmt(off)} seconds ({fmt(off/60)} minutes) off peak. The descriptive gap is {fmt(gap)} seconds ({fmt(gap/60)} minutes), and the ratio is {fmt(ratio,3)}. This supports the amended directional expectation in the observed windows, while the small number of windows and missing route-level Google evidence prevent a significance claim."
    ],("assignment4/tables/period_summary.csv","assignment4/tables/period_summary.json"))
    figure("04_disagreement_metrics.png","Absolute and squared disagreement summaries for every planned window.",[
        "RMSE exceeds MAE in every window, consistent with uneven paired gaps and the greater weight that squared differences give larger gaps. Both measures also remain higher in the peak batches than in the off-peak batches. These summaries strengthen the descriptive pattern across several measures, but they cannot show which routes produced the larger gaps because those Google values were not retained."
    ],("assignment4/tables/batch_results.csv",))
    figure("05_length_comparability.png","Pairs whose returned route lengths differ by at most five percent of their mean.",[
        "The first peak window has 13/20 length-comparable pairs, or 65%, while each other planned window has 15/20, or 75%. Route choice may therefore contribute to the especially large first peak gap. The dashed 70% reference belongs to the original TomTom-2GIS expectation and is shown for context; Google-ORS results do not test that original expectation, and similar lengths do not establish identical geometry."
    ],("assignment4/tables/batch_results.csv",))
    figure("06_ors_routes.png","ORS control durations by directed route across the complete archive.",[
        "All 20 routes have exactly one distinct duration-distance combination across their repeated ORS observations. Repeated requests therefore increased the record count without adding within-route duration variation to this control. In this sample, changing Google-ORS disagreement cannot be interpreted as evidence that both services responded to changing traffic."
    ],("assignment4/tables/route_summary.csv",))
    p("The separate exploratory cycle contains 20 valid pairs, a median absolute gap of 730.50 seconds and 15/20 length-comparable pairs. It occurred outside the defined peak and off-peak categories. Its role is to demonstrate the paired instrument and aggregate calculations, and none of its values enters the period summaries above.")
    evidence("assignment4/tables/batch_results.csv")

    h("3 5 Results by original objective")
    objectives=[
        "Select 20 representative arterial origin-destination routes in Astana and define fixed evening peak and off-peak observation periods.",
        "Collect synchronized Google Maps and Yandex Maps traffic-aware travel-time estimates for the same routes and timestamps over 14 consecutive days.",
        "Calculate absolute travel-time difference, relative percentage difference, MAPE, sMAPE, MAE, RMSE, median absolute error, signed bias, dispersion measures, and 95% confidence intervals for the paired predictions.",
        "Compare the magnitude and variability of API discrepancies between evening peak and off-peak periods using appropriate inferential tests and effect-size measures.",
        "Determine whether discrepancy patterns vary systematically by route, time of day, and traffic period."
    ]
    objective_results=[
        ("20 fixed directed routes; predefined periods. Purposive selection does not establish citywide representativeness.","Figure 6; routes.csv","Partly"),
        ("Original provider pair and 14-day target absent. Amended design has 80 Google-ORS pairs over 2 dates.","Table 1; Table 4","Not achieved as declared"),
        ("Saved batch medians, signed gaps, sMAPE, MAE, RMSE and IQR; no route-level CI, MAPE or complete signed mean.","Table 4; Figures 2-4","Partly"),
        (f"Descriptive contrast {fmt(gap)} s and ratio {fmt(ratio,3)}; no valid route-level inferential test.","Figure 3; Table 8","Partly"),
        ("Window differences shown; systematic route-level Google variation and recurring time patterns remain unresolved.","Figures 2 and 6","Partly")
    ]
    with (TABLES/"objectives.csv").open("w",encoding="utf-8",newline="") as f:
        writer=csv.writer(f,lineterminator="\n");writer.writerow(["objective","declared_wording","result","shown_in","status"])
        for i,(obj,result) in enumerate(zip(objectives,objective_results),1):writer.writerow([f"O{i}",obj,*result])
    table("Results against the five objectives declared in Assignment 1",["ID","Objective as declared","Recorded result and evidence","Status"],
          [[f"O{i}",obj,result[0]+" "+result[1],result[2]] for i,(obj,result) in enumerate(zip(objectives,objective_results),1)],
          [0.8,7,6.6,2.6],sources=("assignment4/tables/objectives.csv",),interpretations=[
              "The amended comparison delivers a narrower result than the original objectives promised. In particular, changing the provider pair and collecting only two dates leaves the original collection objective unmet. Reporting partial completion avoids presenting available descriptive statistics as a substitute for the missing route-level inference."])

    h("3 6 Expected versus observed")
    expected=[]
    for line in (ROOT/"docs/expected_results.md").read_text(encoding="utf-8").splitlines():
        cells=[c.strip() for c in line.strip().strip("|").split("|")]
        if cells and re.fullmatch(r"ER[1-8]",cells[0]):expected.append((cells[0],cells[1]))
    findings=[
        ("0 TomTom-2GIS pairs; Google-ORS is a different comparison.","Not confirmed: untested","Table 1; Figure 3"),
        ("0 eligible original-provider pairs. Google-ORS length shares are 65%, 75%, 75%, 75%.","Not confirmed: untested","Table 1; Figure 5"),
        ("0 original-provider pairs. Positive Google-ORS medians cannot establish this sign.","Not confirmed: untested","Tables 1 and 4"),
        ("0 original-provider pairs. No original peak/off-peak IQR comparison exists.","Not confirmed: untested","Tables 1 and 4"),
        ("0/572 ORS observations fail the recorded valid-route screen (0%).","Confirmed for archived observations","Tables 1-3"),
        ("572/572 ORS requests finish within 10 s (100%); max 5.446 s.","Confirmed for archived observations","Table 3"),
        ("20/20 routes have valid ORS results in all 3 pilot cycles.","Confirmed","Table 1; Figure 6"),
        ("408 automatic attempts out of 768 (53.125%); 4 manual checks excluded. Even all 412 labelled rows fall below 692.","Not confirmed","Table 1; Figure 1")
    ]
    with (TABLES/"expectations.csv").open("w",encoding="utf-8",newline="") as f:
        writer=csv.writer(f,lineterminator="\n");writer.writerow(["id","expected_result_unchanged","observed","verdict","evidence"])
        for item,result in zip(expected,findings):writer.writerow([*item,*result])
    table("Original Assignment 3 expectations copied without changes",["ID","Expected result from Assignment 3","Observed and evidence","Verdict"],
          [[num,text,result[0]+" "+result[2],result[1]] for (num,text),result in zip(expected,findings)],
          [0.8,6.9,6.5,2.8],sources=("docs/expected_results.md","assignment4/tables/expectations.csv","assignment4/tables/inventory.json"),interpretations=[
              "ER8 failed most clearly because the instrument and the shortened collection window did not deliver the original schedule. It is a coverage failure, not evidence about the traffic phenomenon. ER1-ER4 are explicitly untested rather than falsified: access and design changes left their required provider pair unobserved."])
    p("The amended directional expectation stated that peak-window median absolute Google-ORS disagreement would exceed off-peak disagreement. Figure 3 supports that direction for the recorded windows. Its ratio of 1.664 also exceeds the earlier 1.5 descriptive reference, but that does not confirm the original TomTom-2GIS ER1; the provider comparison changed, and uncertainty remains substantial.")
    evidence("docs/google_followup_protocol.json","assignment4/tables/period_summary.json")

    h("3 7 Block E outputs and reliability of collection")
    table("Completed Block E requirement register",["Required output","Actual output and location"],[
        ["Cleaning record","Table 2; archived API and Overpass query, access timestamps and filter rules; zero deleted rows."],
        ["Final dataset","Table 1: 572 retained ORS rows, 100 Google statuses and 5 aggregate batches, of which 4 are planned follow-up."],
        ["Variable list","Table 3 and variables.csv: measurement level, summary and per-variable missingness; overall blanks explicitly defined."],
        ["Trend","Figure 2: dated window medians in seconds/minutes; Figure 1: category counts by collection phase."],
        ["Group comparison","Figure 3: two peak versus two off-peak windows, with contrast and ratio."],
        ["Sensitivity note","Table 8 and window_sensitivity.csv; cleaning-rule and window-deletion checks kept distinct."]
    ],[4,13],sources=("assignment4/evidence_manifest.json","assignment4/tables/variables.csv"),interpretations=[
        "Block E matches the evidence because the services supply the numerical estimates and our scripts sample and analyse them. There are no survey responses, interviews or trained prediction models in the dataset. The completed register therefore focuses on source quality, row accounting, comparisons and sensitivity."])
    p("All 572 collected ORS records enter the retained dataset, and all pass the published screen. The greatest missingness is tied between duration_no_traffic_s and the error message, with 572 blanks in each column; the former is unavailable from ORS and the latter is empty on success. There are zero exact or conflicting duplicate observations after the recorded identity checks, and zero rows outside the duration, distance, timestamp or speed rules. No excluded rows fall disproportionately on a phase or provider because there are no failed ORS exclusions; this does not remove the imbalance in which periods were sampled.")
    evidence("assignment4/tables/inventory.json","assignment4/tables/variables.csv","assignment4/tables/cleaning.csv")
    # This passage separately satisfies the Block E interpretation minimum.
    reliability=[
        "The ORS observation files provide the strongest part of the collection record. Each row has identifiers, a request time, a recorded response status and numerical values. The original response and geometry files remain in the research repository, allowing a reader to inspect how the collector extracted a route. The processed copy adds validation decisions alongside the measurements. This separation makes accidental changes easier to detect and allows the same screening rules to be applied again. Passing those checks establishes internally consistent records; it does not establish that the predicted journey would match a driver on the road.",
        "The route-source evidence is also inspectable. The saved Overpass query and retrieval manifest identify the public source and the time at which the endpoint snapshot was obtained. The route inventory keeps source links for both endpoints. This supports the claim that the routes use fixed public coordinates. It does not establish that every chosen node is a suitable building entrance or that the selected routes represent the journeys made throughout the city. Reverse directions remain separate because road access and routing can differ by direction.",
        "Google evidence has a more restricted form. Request statuses and the comparison protocol support the claim that a batch was attempted at a recorded time, and its counts can be checked against the matching ORS rows. However, the saved aggregate statistics cannot be independently recalculated because individual Google durations and distances were not retained. A transcription or aggregation error in those numeric summaries cannot be ruled out using status records alone. Repeating the collector would obtain new estimates rather than recover the original measurements, so procedural repeatability and exact numerical reproduction must be distinguished.",
        "The batch collection log and GitHub trigger metadata explain another limit: a successful request pipeline can still produce incomplete temporal coverage. The original automatic target was missed even though every retained response passed the checks. The final paired windows repair the missing last batch in the earlier export, but they do not repair the original observation period or turn the small set of dates into a representative time sample. The revised evidence supports a bounded descriptive answer whose reliability is strongest for the archived ORS values and weaker for independently checking the Google aggregates."
    ]
    for text in reliability:p(text)
    evidence("docs/research_log.md","docs/cloud_runs.json","assignment4/evidence_manifest.json","docs/google_followup_protocol.json")

    h("3 8 Sensitivity of the comparison")
    sens=read("sensitivity.csv")
    table("Sensitivity to the stated cleaning choices",["Decision","ORS retained","Valid follow-up pairs","Period gap s","Ratio"],
          [[r["rule"],r["ors_kept"],r["valid_followup_pairs"],fmt(r['peak_minus_offpeak_s']),fmt(r['peak_offpeak_ratio'],3)] for r in sens],
          [6.3,2.4,3,2.8,2.5],sources=("assignment4/tables/sensitivity.csv","assignment4/tables/window_sensitivity.csv"),interpretations=[
              "Allowing low speeds or reducing the upper limit to 160 km/h excludes no ORS observation because all recorded implied speeds lie between 30.01 and 53.80 km/h. Removing the 60 pilot rows leaves 512 ORS observations and the same 80 follow-up pairs. These checks leave the comparison unchanged, but cannot test alternative cleaning of individual Google values because those values are unavailable."])
    p("Window-deletion checks address a different concern: dependence on the few sampled times. Omitting one recorded window at a time leaves a positive peak-off-peak contrast from 202.10 to 340.25 seconds. The corresponding ratios range from 1.495 to 1.833, so the directional pattern persists but the 1.5 descriptive reference is not stable when the first peak window is omitted. These ranges are sensitivity results, not confidence intervals.")
    evidence("assignment4/tables/window_sensitivity.csv")

    h("3 9 Additional interpretation and answer to the question")
    p("The strongest contrast is between the first peak window and the morning off-peak window: their median absolute gaps differ by 369.45 seconds. The broad period contrast remains positive when each window is omitted, yet its magnitude changes noticeably. This makes the direction more defensible within the observed sample than a claim that a specific ratio is a recurring property of Astana travel.")
    p("A non-finding concerns variability. The two peak signed IQRs are 218.30 and 230.23 seconds, below the off-peak values of 261.30 and 279.48 seconds. Thus larger central disagreement was not accompanied by a larger within-window signed spread in this amended comparison. This observation does not test original ER4 because that expectation concerns TomTom and 2GIS.")
    p("Provider-specific road speeds and route selection could produce the same broad pattern as different responses to congestion. The constant ORS estimates demonstrate that this control contributed no observed within-route temporal change. The lower length-comparability share in the first peak batch supplies a further route-choice explanation, but without retained Google geometry its contribution cannot be separated from other provider differences.")
    p("The scope is the fixed public-endpoint car routes and the recorded windows on the two study dates. The results do not cover buses, walking, every Astana district, weekday commuting across a longer period, or the original provider pair. Both dates fall on a weekend, which further limits any claim about typical working-day congestion.")
    p("The most useful surprise is the complete stability of ORS durations and distances within each route. This means that a large archive of requests can still contain little temporal information about one provider. It also explains why record counts and the effective amount of independent comparison evidence must be reported separately.")
    p(f"Answer to the research question: on the sampled Astana routes and dates, Google-ORS median absolute disagreement was greater in the recorded peak windows, with a median-of-window-medians gap of {fmt(gap)} seconds relative to off peak (Figure 3), without establishing accuracy or a causal traffic effect.")
    evidence("assignment4/tables/batch_results.csv","assignment4/tables/route_summary.csv","assignment4/tables/period_summary.json","assignment4/tables/reported_metrics.json")

    h("3 10 Comparison with the Assignment 2 literature")
    literature=[
        ["Derrow-Pinion et al. (2021)","Road structure and traffic inputs inform Google ETA prediction.","The gap varies across windows while ORS remains constant.","Consistent with the motivation; does not test its model","Our outputs do not identify Google's current internal model."],
        ["Soon et al. (2023)","Forecasting depends on temporal scope, data quality and unusual conditions.","Only two dates and four paired windows are available.","Supports the scope caution","The current sample cannot establish recurring temporal patterns."],
        ["Fu et al. (2023)","Common endpoints give different estimates across routing methods; Google is a benchmark.","Google and ORS differ by several minutes in all windows.","Extends to sampled Astana routes","Disagreement remains visible with fixed endpoints; neither comparator is observed truth."],
        ["Bastarianto et al. (2025)","API and reported journeys disagree, with preprocessing and mode affecting the result.","Our provider gap is measurable, but actual car journeys are absent.","Supports caution; numerical findings not comparable","A provider difference cannot estimate the reported/API mismatch studied there."],
        ["Boeing and Zhou (2026)","An open-data model targets minimally congested travel and is assessed against Google.","The ORS-control gap is larger in our peak summaries.","Extends the question to period dependence","Our comparison does not evaluate their trained model or its prediction quality."],
        ["Mansurova et al. (2025)","Astana bus GPS records distinguish movement and dwell behaviour.","Our car API outputs do not measure either bus movement or dwell.","Not directly comparable","Local geography is shared, but mode and measurement differ; bus evidence cannot validate our car estimates."],
        ["Hodson (2022)","Absolute and squared losses serve different evaluation aims.","RMSE exceeds MAE in every archived comparison window.","Supports distinct metric purposes","The squared summary highlights larger gaps; neither summary supplies real journey accuracy."],
        ["Valero et al. (2025)","Calculated and reported commute times diverge for car and transit journeys.","We establish disagreement only between provider outputs.","Supports the reference-data caution","No reported or instrumented trip was collected, so their mismatch cannot be confirmed here."],
        ["Elnaggar et al. (2025)","Repeated API sampling reveals temporal and route differences; reported validation shows close agreement.","Peak summaries exceed off-peak summaries, but our control gap remains substantial.","Supports temporal sampling; diverges in reported closeness","Comparator, city and sampling differ; navigation-estimate agreement is a different validation target."],
        ["Verma et al. (2025)","Routing outputs can have high statistical fit while differing in average duration.","Window gaps persist; no correlation is used as an agreement claim.","Supports difference-based reporting","Minutes of disagreement answer our question more directly than association alone."]
    ]
    table("Relation to the ten sources retained in Assignment 2",["Source","Finding","Our result","Relation","Reason"],literature,
          [3.1,3.5,3.5,3.3,3.6],sources=("assignment4/tables/batch_results.csv","assignment4/tables/period_summary.csv"),interpretations=[
              "The literature comparison shows where the studies reach different kinds of findings rather than claiming universal agreement. Close agreement in the Jeddah validation and persistent mismatch in reported-journey studies concern different references, while our study compares two routing outputs. We therefore extend evidence on provider disagreement without using it to confirm or contradict a numerical accuracy result from another city."])
    p("The Assignment 2 review retained ten papers, so all ten are shown rather than selecting only those with a convenient result. The most direct empirical connections are the routing comparisons and repeated temporal sampling. The local bus dataset remains relevant to context and measurement design, with its lack of direct comparability stated openly.")

    h("3 11 Validity and limitations")
    table("Specific limits on the answer",["Validity issue","What in this study creates the limit","Consequence"],[
        ["Internal interpretation","Google uses traffic information; ORS does not, and paths/speed models can differ.","A peak difference cannot isolate the effect of traffic."],
        ["Measurement validity","No independently observed car journeys are available.","No provider accuracy ranking or congestion-delay estimate is established."],
        ["Reliability","Google durations and geometry are not archived individually.","Saved Google numerical summaries cannot be independently recalculated."],
        ["Temporal generalisation","Two weekend dates, four short comparison batches and a shortened final window.","Working-day and longer-term patterns remain unresolved."],
        ["Sampling and dependence","Purposive routes share endpoints and recur across windows.","80 pairs are not 80 independent route/day samples; no citywide inference."],
        ["Original objectives","Original providers, 14-day period and inferential dataset were not delivered.","Earlier objectives remain partly or wholly unmet despite the amended result."],
        ["Coverage and retention","Original automatic target missed; comparable Google geometry absent.","Missing intervals and route-level explanations cannot be repaired retrospectively."]
    ],[3.4,7.2,6.4],sources=("docs/google_followup_protocol.json","assignment4/tables/inventory.json","assignment4/tables/route_summary.csv"),interpretations=[
              "These limits change what the result can answer, rather than merely reducing confidence in an otherwise general claim. The data establish a descriptive difference for the observed windows while leaving accuracy, traffic causation and recurring citywide patterns open. Additional requests using the same aggregate-only retention would increase counts without resolving every one of these limits."])

    segment="conclusion"
    h("4 Conclusion",1)
    for text in [
        f"Google and ORS disagreed more in the recorded peak windows than in the recorded off-peak windows on the sampled Astana routes. The median of window median absolute differences was {fmt(peak)} seconds at peak and {fmt(off)} seconds off peak, giving a descriptive contrast of {fmt(gap)} seconds. This answers the amended research question for the observed routes and dates. It does not show that either service predicts actual driving time more accurately or that traffic alone caused the difference.",
        "The contribution is a documented comparison with clear boundaries. All planned comparison windows are represented, the collection phases remain separate, and the central differences are accompanied by spread and route-length evidence. The larger peak disagreement is visible across the recorded summaries. Its exact size remains sensitive to which windows are retained, so the observed ratio should not become a general adjustment factor for navigation estimates in Astana.",
        "The control result is especially important for understanding the evidence. ORS returned unchanged duration and distance values within every route. Repeating these requests produced more records but no observed temporal variation in the control estimate. The provider gap therefore combines the behaviour of a changing traffic-aware output with a stable comparison. Different routing choices and speed assumptions remain possible explanations alongside changing traffic conditions.",
        "The study also leaves original commitments unresolved. The earlier provider comparison was not performed, the original automatic collection target was missed, and the original route-level inferential analysis could not be supported. The Google aggregates cannot be independently recalculated from saved individual responses. These are substantive limits on the claims available from the project, even though the archived ORS observations passed the recorded checks. A successful amended comparison does not erase the unmet objectives or test the original provider-specific expectations.",
        "Future work should first obtain permission to retain auditable paired numerical outputs and both route geometries. It should then cover working days and weekends across a longer observation period, with repeated matched routes in each period. Independently timed car journeys on those same routes would make an accuracy assessment possible. Until such evidence exists, the present findings are useful as a local example of provider disagreement and as a warning that agreement with another routing service cannot replace validation against travel on the road."
    ]:p(text)

    segment="references"
    h("References",1)
    refs=[
        ("Bastarianto, F.F., Hancock, T.O., Ilahi, A., Manley, E. and Choudhury, C.F. (2025) 'Mind the Gap: The impact of discrepancies between Google Maps API and reported travel data in the Global South', Case Studies on Transport Policy, 21, 101508. doi:10.1016/j.cstp.2025.101508.","https://eprints.whiterose.ac.uk/id/eprint/228392/1/1-s2.0-S2213624X25001452-main.pdf"),
        ("Boeing, G. and Zhou, Y. (2026) 'Travel time prediction from sparse open data', International Journal of Geographical Information Science, 40(10), pp. 3412-3428. doi:10.1080/13658816.2026.2628193.","https://arxiv.org/pdf/2602.15069"),
        ("Derrow-Pinion, A., She, J., Wong, D., Lange, O., Hester, T., Perez, L., Nunkesser, M., Lee, S., Guo, X., Wiltshire, B., Battaglia, P.W., Gupta, V., Li, A., Xu, Z., Sanchez-Gonzalez, A., Li, Y. and Velickovic, P. (2021) 'ETA Prediction with Graph Neural Networks in Google Maps', Proceedings of the 30th ACM International Conference on Information and Knowledge Management, pp. 3767-3776. doi:10.1145/3459637.3481916.","https://luisperez.me/files/eta-prediction.pdf"),
        ("Elnaggar, G.R., Al-Hourani, S. and Abutaha, R. (2025) 'Real-Time Urban Congestion Monitoring in Jeddah, Saudi Arabia, Using the Google Maps API: A Data-Driven Framework for Middle Eastern Cities', Sustainability, 17(18), 8194. doi:10.3390/su17188194.","https://mdpi-res.com/d_attachment/sustainability/sustainability-17-08194/article_deploy/sustainability-17-08194.pdf"),
        ("Fu, X., Kakkar, D., Chen, J., Moynihan, K.M., Hegland, T.A. and Blossom, J. (2023) 'A comparative study of methods for drive time estimation on geospatial big data: A case study in USA', The International Archives of the Photogrammetry, Remote Sensing and Spatial Information Sciences, XLVIII-4/W7-2023, pp. 53-60. doi:10.5194/isprs-archives-XLVIII-4-W7-2023-53-2023.","https://isprs-archives.copernicus.org/articles/XLVIII-4-W7-2023/53/2023/isprs-archives-XLVIII-4-W7-2023-53-2023.pdf"),
        ("Hodson, T.O. (2022) 'Root-mean-square error (RMSE) or mean absolute error (MAE): when to use them or not', Geoscientific Model Development, 15(14), pp. 5481-5487. doi:10.5194/gmd-15-5481-2022.","https://gmd.copernicus.org/articles/15/5481/2022/gmd-15-5481-2022.pdf"),
        ("Mansurova, A., Mussina, A., Aubakirov, S., Nugumanova, A. and Yedilkhan, D. (2025) 'From Raw GPS to GTFS: A Real-World Open Dataset for Bus Travel Time Prediction', Data, 10(8), 119. doi:10.3390/data10080119.","https://mdpi-res.com/d_attachment/data/data-10-00119/article_deploy/data-10-00119.pdf"),
        ("Soon, K.L., Chan, R.K.C., Lim, J.M.-Y. and Parthiban, R. (2023) 'Short-term traffic forecasting model: prevailing trends and guidelines', Transportation Safety and Environment, 5(3), tdac058. doi:10.1093/tse/tdac058.","https://researchmgt.monash.edu/ws/portalfiles/portal/528840794/488448406_oa.pdf"),
        ("Valero, A., Negm, H. and El-Geneidy, A. (2025) 'Exploring the (Mis)match between Reported and Calculated Commute Travel Times', Findings, August. doi:10.32866/001c.143380.","https://findingspress.org/article/143380.pdf"),
        ("Verma, R., Mittal, S. and Ukkusuri, S.V. (2025) 'Spatial Access of America: Multiple indicators of accessibility to opportunities', Scientific Data, 12, 1223. doi:10.1038/s41597-025-05440-8.","https://www.nature.com/articles/s41597-025-05440-8.pdf")
    ]
    for text,url in refs:
        para=p(text+" Full text: ");link(para,"Open article",url)
    for name,url in [
        ("Google (n.d.) 'Configure routing preferences'. Google Maps Platform documentation.","https://developers.google.com/maps/documentation/routes/config_trade_offs"),
        ("HeiGIT (n.d.) 'openrouteservice API reference'.","https://giscience.github.io/openrouteservice/api-reference/"),
        ("OpenStreetMap contributors (n.d.) 'Copyright and licence'.","https://www.openstreetmap.org/copyright")
    ]:
        para=p(name+" Available at: ");link(para,"Source documentation",url);para.add_run(" (Accessed: 31 October 2026).")
    h("Appendix A Evidence and reproducibility",1)
    p("The public data repository contains the submission evidence under research: collection scripts, source records, request settings, a batch log derived from the archived observations, analysis tables and figures. It is a publication copy of the collection evidence, not an independent collection source. Published request-setting files identify their projection from the historical protocols; the original collection archive is preserved. The final snapshot in research/assignment4/data fixes the report inputs.")
    para=p("Browsable evidence: ");link(para,"Assignment 4 folder",BASE+"/tree/main/research/assignment4");para.add_run("; ");link(para,"Data export",DATA_BASE+"/tree/main")
    p("From the published research folder, run python assignment4/analyse.py with the declared plotting dependencies, then python assignment4/build_report.py with python-docx. The evidence manifest records source hashes, generated table hashes and figure descriptions. Requirements-used.txt records the exact analysis environment; the submission-checks file gives the counted text, visuals and unchanged expectations.")
    evidence("assignment4/evidence_manifest.json","assignment4/submission_checks.json","assignment4/requirements-used.txt")
    p("ORS outputs: © openrouteservice by HeiGIT | Data from OpenStreetMap, under the applicable CC BY-SA terms. Endpoint source: © OpenStreetMap contributors under ODbL. Google individual numerical outputs and geometry are absent; only the recorded comparison aggregates and request metadata are supplied.")
    h("Appendix B Instrument provenance and original scenarios",1)
    table("Instrument and source register",["Item","Provenance and role"],[
        ["API collector","Custom Python collector in collector/; fixed coordinates, recorded request outcomes and bounded request budgets."],
        ["Repeated sampling","Adapted from Elnaggar et al. (2025); applied to the recorded Astana windows, without claiming the same temporal coverage."],
        ["Common endpoint comparison","Adapted from Fu et al. (2023); provider choice and Google retention were amended explicitly."],
        ["Disagreement metrics","Informed by Hodson (2022); differences interpreted as agreement quantities, with no measured-trip reference."],
        ["Assignment 4 analysis","Offline script assignment4/analyse.py; reads archived evidence and creates auditable tables, sensitivity checks and figures."]
    ],[4,13],numbered=False)
    p("The instrument is an original implementation informed by published approaches, rather than a replication of a single paper. New Assignment 4 processing is distinguishable from the scripts that made the live requests.")
    scenarios=[
        ["bad","Traffic-provider permissions are still unavailable; only control-instrument evidence exists and H1 cannot be tested."],
        ["bad","At least 2% of ORS requests fail or fewer than 20 routes work; revise the instrument in a dated amendment."],
        ["medium","All routes work but fewer than 90% of scheduled attempts are recorded; keep the actual times and report the incomplete coverage."],
        ["medium","Primary collection later becomes permitted but period contrast is below 1.5; report evidence against ER1 without editing it."],
        ["good","The ORS pilot covers all routes and fewer than 2% of requests fail; the data pipeline is ready for a permitted primary comparison."],
        ["good","A later permitted primary sample spans both periods and reaches the ER1 threshold with interpretable route-aware uncertainty."]
    ]
    table("Method-related scenarios preserved from Assignment 3",["Outcome","Original wording"],scenarios,[2.5,14.5],numbered=False)
    p("These scenarios retain the original design's conditions and are not rewritten to match the amended providers. Table 6 gives the explicit verdicts for the original numbered expectations, including the observed coverage failure.")

    results_words=sum(words(s) for s in body)
    conclusion_words=sum(words(s) for s in conclusion_text)
    count_para.text=f"Results and Discussion only: {results_words} words; {figure_count} figures/graphs; {table_count} tables. Table text, captions, headings and evidence-link lines are excluded from the word count. Conclusion: {conclusion_words} words."
    foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");foot._p.append(field)
    doc.save(OUTPUT)
    counts=dict(abstract_words=words(abstract),results_discussion_words=results_words,conclusion_words=conclusion_words,
                block_e_reliability_words=sum(words(s) for s in reliability),results_figures=figure_count,results_tables=table_count,
                rq=RQ,original_expectation_rows=len(expected),table_interpretations=table_interpretations,
                figure_interpretations=figure_interpretations,report=OUTPUT.relative_to(ROOT).as_posix(),
                report_sha256=hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
                word_count_rule="Body paragraphs only in Results and Discussion; excludes tables, captions, headings and evidence links.")
    assert counts["abstract_words"]<=150 and results_words>=800 and conclusion_words>=300
    assert counts["block_e_reliability_words"]>=300 and figure_count>=5 and table_count>=3
    (HERE/"submission_checks.json").write_text(json.dumps(counts,indent=2),encoding="utf-8")
    (HERE/"report_links.json").write_text(json.dumps(sorted(links),indent=2),encoding="utf-8")
    (HERE/"report_text.txt").write_text("\n\n".join(p.text for p in doc.paragraphs),encoding="utf-8")
    metric_register=dict(
        peak_summary_s=peak,offpeak_summary_s=off,contrast_s=gap,ratio=ratio,
        peak_summary_minutes=peak/60,offpeak_summary_minutes=off/60,contrast_minutes=gap/60,
        largest_window_median_contrast_s=max(float(r["median_absolute_difference_s"]) for r in follow)-min(float(r["median_absolute_difference_s"]) for r in follow),
        rounded_values=dict(peak_s=fmt(peak),offpeak_s=fmt(off),contrast_s=fmt(gap),ratio=fmt(ratio,3),
                            peak_min=fmt(peak/60),offpeak_min=fmt(off/60),contrast_min=fmt(gap/60)),
        definitions="Period values are medians of window median absolute differences. Contrast subtracts off-peak from peak. Ratio divides peak by off-peak. Largest contrast subtracts the smallest follow-up window median from the largest.")
    (TABLES/"reported_metrics.json").write_text(json.dumps(metric_register,indent=2),encoding="utf-8")
    manifest=json.loads((HERE/"evidence_manifest.json").read_text())
    for name in ["objectives.csv","expectations.csv","reported_metrics.json"]:
        file=TABLES/name
        manifest["outputs"][file.relative_to(ROOT).as_posix()]=text_hash(file)
    manifest["scripts"]["assignment4/build_report.py"]=text_hash(Path(__file__))
    (HERE/"evidence_manifest.json").write_text(json.dumps(manifest,indent=2),encoding="utf-8")
    print(json.dumps({k:v for k,v in counts.items() if k in ["abstract_words","results_discussion_words","conclusion_words","block_e_reliability_words","results_figures","results_tables"]}))


if __name__ == "__main__":
    build()
